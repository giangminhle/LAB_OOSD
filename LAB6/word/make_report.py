# -*- coding: utf-8 -*-
"""Tạo báo cáo Word LAB6 ngắn gọn, có ảnh sơ đồ và ảnh chụp form."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIAG = os.path.join(BASE, 'diagrams')
WORD = os.path.join(BASE, 'word')
OUT = os.path.join(WORD, 'BaoCao_LAB6_QuanLyBenhAn.docx')

FONT = 'Times New Roman'


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def p(text='', size=12, bold=False, italic=False, align=None, space_after=4,
      space_before=0, color=None, first_line=None):
    par = doc.add_paragraph()
    if align is not None:
        par.alignment = align
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.line_spacing = 1.15
    if first_line:
        par.paragraph_format.first_line_indent = Cm(first_line)
    run = par.add_run(text)
    set_run_font(run, size, bold, italic, color)
    return par


def heading(text, level=1):
    par = doc.add_paragraph()
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    par.paragraph_format.space_after = Pt(5)
    run = par.add_run(text)
    set_run_font(run, 14 if level == 1 else 12.5, True, False,
                 (0x1F, 0x3B, 0x63) if level == 1 else (0, 0, 0))
    # gắn outline level để Word hiện trong Navigation
    pPr = par._p.get_or_add_pPr()
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), str(level - 1))
    pPr.append(outline)
    return par


def bullet(text, size=12):
    par = doc.add_paragraph(style='List Bullet')
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.15
    run = par.add_run(text)
    set_run_font(run, size)
    return par


def caption(text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(8)
    run = par.add_run(text)
    set_run_font(run, 10, False, True, (0x40, 0x40, 0x40))


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def cell_margins(cell, top=40, bottom=40, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tcPr.append(mar)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def add_table(headers, rows, widths, font_size=9, header_bg='D9E2F3', align_center=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), str(int(sum(widths) * 1440)))
    tblPr.append(tblW)
    for el in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(el)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    hdr = table.rows[0]
    repeat_header(hdr)
    cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        shade(cell, header_bg)
        cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        run = par.add_run(h)
        set_run_font(run, font_size, True)
    for row_data in rows:
        row = table.add_row()
        cant_split(row)
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            if align_center and i in align_center:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(str(val))
            set_run_font(run, font_size)
    return table


def add_picture_center(path, width_in):
    # giữ đoạn mô tả phía trên và chú thích phía dưới cùng trang với hình
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    par = doc.add_paragraph()
    par.paragraph_format.keep_with_next = True
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run()
    run.add_picture(path, width=Inches(width_in))
    return par


def add_page_number(par):
    run = par.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    set_run_font(run, 9)


def set_section_a4(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.6)


def header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run('LAB6 - Quản lý hồ sơ bệnh án')
    set_run_font(r, 9, italic=True, color=(0x59, 0x59, 0x59))
    fp = section.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('Trang ')
    set_run_font(r, 9)
    add_page_number(fp)


doc = Document()
style = doc.styles['Normal']
style.font.name = FONT
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
set_section_a4(doc.sections[0])
header_footer(doc.sections[0])

# ============================== BÌA ==============================
for _ in range(3):
    p('', space_after=0)
p('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM (HUTECH)', 14, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('KHOA CÔNG NGHỆ THÔNG TIN', 13, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('BÁO CÁO THỰC HÀNH', 18, True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=30)
p('PHÂN TÍCH THIẾT KẾ HƯỚNG ĐỐI TƯỢNG', 14, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('LAB 6 (BÀI 5)', 22, True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=30)
p('QUẢN LÝ HỒ SƠ BỆNH ÁN CỦA BỆNH NHÂN', 18, True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)
info = [
    ('Họ và tên:', 'Lê Minh Giang'),
    ('MSSV:', '1250080043'),
    ('Lớp:', '12_DH_CNPM1'),
    ('Nhóm TH:', '................................'),
    ('Giảng viên hướng dẫn:', '................................'),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
no_borders(t)
for k, v in info:
    row = t.add_row()
    row.cells[0].width = Inches(2.2)
    row.cells[1].width = Inches(3.5)
    pa = row.cells[0].paragraphs[0]
    pa.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(pa.add_run(k), 12, True)
    set_run_font(row.cells[1].paragraphs[0].add_run(v), 12)
doc.add_page_break()

# ============================== 1. MÔ TẢ ==============================
heading('1. MÔ TẢ BÀI TOÁN', 1)
p('Bệnh viện đa khoa quốc tế cần phần mềm quản lý hồ sơ bệnh án. Bệnh nhân khám lần đầu sẽ điền phiếu '
  'đăng ký (họ tên, địa chỉ, ngày sinh, điện thoại, ảnh 3x4) và đóng lệ phí năm, sau đó được cấp hồ sơ '
  'bệnh án có mã hồ sơ riêng. Mỗi lần khám, bác sĩ ghi phiếu khám (ngày giờ, khoa, triệu chứng, thân nhiệt, '
  'huyết áp), chỉ định xét nghiệm nếu cần và kê toa thuốc. Bệnh nhân thanh toán viện phí xong thì phòng '
  'chăm sóc khách hàng ghi nhận và hoàn tất đợt điều trị. Lần sau chỉ cần báo mã hồ sơ hoặc họ tên, số '
  'điện thoại là tra được lịch sử khám. Bệnh nhân có hồ sơ được khám định kỳ theo quý miễn phí.',
  first_line=1.0)

# ============================== 2. BÀI LÀM ==============================
heading('2. BÀI LÀM', 1)
heading('Câu 1. Xác định Class, thuộc tính, phương thức và ID', 2)
p('Em xác định được 13 class, mỗi class có một thuộc tính ID làm khóa chính:', first_line=1.0)
classes = [
    ('1', 'BenhNhan', 'MaBN', 'HoTen, NgaySinh, DiaChi, DienThoai, Anh3x4', 'DangKyLanDau()'),
    ('2', 'PhieuDangKy', 'MaPhieuDK', 'MaBN, NgayDK, HoTen, DiaChi, NgaySinh, DienThoai', 'LapPhieu(), XacNhan()'),
    ('3', 'LePhiNam', 'MaLePhi', 'MaBN, Nam, SoTien, NgayDong, NguoiThu', 'DongLePhi()'),
    ('4', 'HoSoBenhAn', 'MaHoSo', 'MaBN, MaHSCode, NgayCap, TrangThai', 'CapHoSo(), TraCuu()'),
    ('5', 'LichHenKham', 'MaLichHen', 'MaHoSo, NgayGioHen, LoaiHen, TrangThai', 'DatLichHen(), HuyLichHen()'),
    ('6', 'Khoa', 'MaKhoa', 'TenKhoa', 'ThemKhoa()'),
    ('7', 'BacSi', 'MaBS', 'HoTen, MaKhoa, DienThoai', 'KhamBenh(), KeToa()'),
    ('8', 'PhieuKham', 'MaPhieuKham', 'MaHoSo, MaBS, MaKhoa, NgayGioKham, TrieuChung, ThanNhiet, HuyetAp, '
                                       'ChanDoan, LoaiKham', 'TaoPhieuKham(), HoanTat()'),
    ('9', 'XetNghiem', 'MaXN', 'MaPhieuKham, LoaiXN, NgayChiDinh, KetQua', 'ChiDinh(), NhapKetQua()'),
    ('10', 'ToaThuoc', 'MaToa', 'MaPhieuKham, MaBS, NgayKe, GhiChu', 'LapToa()'),
    ('11', 'ChiTietToa', 'MaToa + MaThuoc', 'LieuLuong, SoLuong, CachDung', 'ThemThuoc()'),
    ('12', 'Thuoc', 'MaThuoc', 'TenThuoc, DonVi, DonGia', 'CapNhatGia()'),
    ('13', 'HoaDonVienPhi', 'MaHD', 'MaPhieuKham, NgayLap, TongTien, DaThanhToan', 'ThanhToan()'),
]
add_table(['STT', 'Class', 'ID', 'Thuộc tính', 'Phương thức'], classes,
          [0.4, 1.3, 1.1, 2.65, 1.5], font_size=9, align_center=[0])

heading('Câu 2. Liên kết giữa các Class và Class Diagram', 2)
assoc = [
    ('BenhNhan - PhieuDangKy', '1 - 1..*', 'Bệnh nhân có ít nhất 1 phiếu đăng ký.'),
    ('BenhNhan - LePhiNam', '1 - 1..*', 'Mỗi năm đóng 1 lần lệ phí.'),
    ('BenhNhan - HoSoBenhAn', '1 - 1', 'Mỗi bệnh nhân chỉ có 1 hồ sơ.'),
    ('HoSoBenhAn - LichHenKham', '1 - 0..*', 'Hẹn khám định kỳ theo quý.'),
    ('HoSoBenhAn - PhieuKham', '1 - 0..*', 'Hồ sơ ghi nhiều lần khám.'),
    ('Khoa - BacSi', '1 - 1..*', 'Bác sĩ thuộc 1 khoa.'),
    ('Khoa - PhieuKham', '1 - 0..*', 'Phiếu khám ghi khoa điều trị.'),
    ('BacSi - PhieuKham', '1 - 0..*', 'Bác sĩ khám nhiều phiếu.'),
    ('PhieuKham - XetNghiem', '1 - 0..*', 'Có thể không cần xét nghiệm.'),
    ('PhieuKham - ToaThuoc', '1 - 0..1', 'Mỗi phiếu khám tối đa 1 toa.'),
    ('PhieuKham - HoaDonVienPhi', '1 - 0..1', 'Mỗi phiếu khám có 1 hóa đơn.'),
    ('ToaThuoc - ChiTietToa', '1 - 1..*', 'Toa có ít nhất 1 thuốc.'),
    ('Thuoc - ChiTietToa', '1 - 0..*', 'Thuốc xuất hiện trong nhiều toa.'),
]
add_table(['Liên kết', 'Bản số', 'Giải thích'], assoc, [2.4, 0.9, 3.65], font_size=9, align_center=[1])
add_picture_center(os.path.join(DIAG, '1_ClassDiagram.drawio.png'), 6.9)
caption('Hình 1 - Class Diagram')

heading('Câu 3. Actor và Use Case tổng quát', 2)
p('Actor: Bệnh nhân, NV chăm sóc khách hàng, NV thu ngân, Bác sĩ, Quản trị viên.', first_line=1.0)
p('Use case: Đăng ký khám lần đầu, Đóng lệ phí năm, Cấp hồ sơ bệnh án, Đặt lịch hẹn định kỳ quý, '
  'Tiếp nhận và tra cứu hồ sơ, Khám bệnh, Chỉ định xét nghiệm, Lập toa thuốc, Thanh toán viện phí, '
  'Hoàn tất đợt điều trị, Tra cứu lịch sử và báo cáo, Quản lý danh mục.', first_line=1.0)
add_picture_center(os.path.join(DIAG, '2_UseCase_TongQuat.drawio.png'), 5.6)
caption('Hình 2 - Use Case Diagram tổng quát')

heading('Câu 4. Use Case phân rã: đăng ký khám - xét nghiệm - nhận toa thuốc', 2)
p('Em phân rã quy trình thành 10 use case theo thứ tự: nhập phiếu đăng ký, đóng lệ phí năm, cấp hồ sơ, '
  'tra cứu hồ sơ, tạo phiếu khám, chỉ định và nhập kết quả xét nghiệm, lập toa thuốc, thanh toán viện phí, '
  'nhận toa và hoàn tất. Chỉ định xét nghiệm là «extend» của tạo phiếu khám vì chỉ làm khi cần.',
  first_line=1.0)
add_picture_center(os.path.join(DIAG, '3_UseCase_PhanRa_DangKy_Toa.drawio.png'), 4.6)
caption('Hình 3 - Use Case Diagram phân rã')

heading('Câu 5. Activity Diagram', 2)
p('Sơ đồ có 3 điều kiện rẽ nhánh: thông tin đăng ký có hợp lệ không, có cần xét nghiệm không và đã '
  'thanh toán viện phí chưa.', first_line=1.0)
add_picture_center(os.path.join(DIAG, '4_Activity_DangKy_Toa.drawio.png'), 4.0)
caption('Hình 4 - Activity Diagram đăng ký khám đến nhận toa thuốc')

heading('Câu 6. Sequence Diagram “Lập toa thuốc”', 2)
p('Bác sĩ chọn phiếu khám và nhập thuốc trên FrmKhamBenh. Form gọi KhamBenhService.LapToaThuoc(), service '
  'kiểm tra dữ liệu, nếu hợp lệ thì mở transaction: thêm ToaThuoc, thêm từng dòng ChiTietToa, cập nhật phiếu '
  'khám sang “Đã khám” rồi Commit. Nếu lỗi thì Rollback và báo lỗi cho bác sĩ.', first_line=1.0)
add_picture_center(os.path.join(DIAG, '5_Sequence_LapToaThuoc.drawio.png'), 5.8)
caption('Hình 5 - Sequence Diagram lập toa thuốc')

# ============================== 3. CHƯƠNG TRÌNH ==============================
heading('3. CHƯƠNG TRÌNH DEMO', 1)
p('Em làm ứng dụng C# WinForms kết nối SQL Server (CSDL QuanLyBenhAn), chia 3 lớp UI - Services - Data. '
  'Chương trình có 2 chức năng chính và 1 nút chạy kiểm thử quy tắc nghiệp vụ.', first_line=1.0)
shots = [
    ('Screenshot 2026-10-16 105019.png', 'Hình 6 - Form đăng ký lần đầu, đóng lệ phí, cấp hồ sơ', 4.0),
    ('Screenshot 2026-10-16 105029.png', 'Hình 7 - Form khám bệnh, xét nghiệm, toa thuốc, thanh toán', 6.5),
    ('Screenshot 2026-10-16 104639.png', 'Hình 8 - Form chính và kết quả kiểm thử 16/16 PASS', 4.6),
]
for fn, cap, w in shots:
    add_picture_center(os.path.join(WORD, fn), w)
    caption(cap)

tests = [
    ('DK01 - DK03', 'Đăng ký hợp lệ, chặn SĐT sai, chặn ngày sinh tương lai', 'PASS'),
    ('LP01 - LP02', 'Lệ phí năm hợp lệ, chặn số tiền bằng 0', 'PASS'),
    ('Q01 - Q02', 'Lịch định kỳ cách nhau 1 quý, chặn hẹn cách 2 tuần', 'PASS'),
    ('K01 - K02', 'Phiếu khám hợp lệ, chặn thân nhiệt 50 độ', 'PASS'),
    ('XN01', 'Xét nghiệm phải có mã phiếu khám', 'PASS'),
    ('T01 - T03', 'Toa phải có thuốc, phải gắn phiếu khám, toa hợp lệ', 'PASS'),
    ('HT01 - HT03', 'Chỉ hoàn tất khi đã có hóa đơn và đã thanh toán', 'PASS'),
]
add_table(['Mã test', 'Nội dung', 'Kết quả'], tests, [1.2, 4.9, 0.85], font_size=9.5, align_center=[0, 2])
caption('Bảng kết quả kiểm thử (16/16 PASS)')

# ============================== 4. KẾT LUẬN ==============================
heading('4. KẾT LUẬN', 1)
p('Em đã làm xong 6 câu của bài: xác định 13 class và ID, vẽ class diagram có bản số, use case tổng quát, '
  'use case phân rã, activity diagram và sequence diagram lập toa thuốc. Chương trình demo chạy được các '
  'nghiệp vụ đăng ký, khám bệnh, kê toa, thanh toán và hoàn tất. Hạn chế là chưa có đăng nhập phân quyền '
  'và chưa in được hồ sơ, toa thuốc ra file.', first_line=1.0)

doc.save(OUT)
print('saved', OUT)


