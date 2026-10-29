# -*- coding: utf-8 -*-
"""Tạo báo cáo Word LAB7 ngắn gọn, có ảnh sơ đồ và ảnh chụp form."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = r'C:\Users\ASUS\OneDrive\NAM4\PTPMHTD\lab\LAB_OOSD\LAB7'
DIAG = os.path.join(BASE, 'diagrams')
WORD = os.path.join(BASE, 'word')
OUT = os.path.join(WORD, 'BaoCao_LAB7_QuanLyNhanSu.docx')

FONT = 'Times New Roman'


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), FONT)


def p(text='', size=12, bold=False, italic=False, align=None, space_after=4,
      space_before=0, color=None, first_line=None):
    par = doc.add_paragraph()
    if align is not None:
        par.alignment = align
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.line_spacing = 1.15
    if first_line:
        par.paragraph_format.first_line_indent = Cm(first_line)
    run = par.add_run(text)
    set_run_font(run, size, bold, italic, color)
    return par


def heading(text, level=1):
    par = doc.add_paragraph()
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    par.paragraph_format.space_after = Pt(5)
    run = par.add_run(text)
    set_run_font(run, 14 if level == 1 else 12.5, True, False,
                 (0x1F, 0x3B, 0x63) if level == 1 else (0, 0, 0))
    # gắn outline level để Word hiện trong Navigation
    pPr = par._p.get_or_add_pPr()
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), str(level - 1))
    pPr.append(outline)
    return par


def bullet(text, size=12):
    par = doc.add_paragraph(style='List Bullet')
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.15
    run = par.add_run(text)
    set_run_font(run, size)
    return par


def caption(text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(8)
    run = par.add_run(text)
    set_run_font(run, 10, False, True, (0x40, 0x40, 0x40))


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def cell_margins(cell, top=40, bottom=40, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tcPr.append(mar)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'nil')
        borders.append(el)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def add_table(headers, rows, widths, font_size=9, header_bg='D9E2F3', align_center=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), str(int(sum(widths) * 1440)))
    tblPr.append(tblW)
    for el in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(el)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    hdr = table.rows[0]
    repeat_header(hdr)
    cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        shade(cell, header_bg)
        cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        run = par.add_run(h)
        set_run_font(run, font_size, True)
    for row_data in rows:
        row = table.add_row()
        cant_split(row)
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            if align_center and i in align_center:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(str(val))
            set_run_font(run, font_size)
    return table


def add_picture_center(path, width_in):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run()
    run.add_picture(path, width=Inches(width_in))
    return par


def add_page_number(par):
    run = par.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar)
    run._r.append(instrText)
    run._r.append(fldChar2)
    set_run_font(run, 9)


def set_section_a4(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.6)


def header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run('LAB7 - Quản lý nhân sự - Công ty Sơn Đông Á')
    set_run_font(r, 9, italic=True, color=(0x59, 0x59, 0x59))
    fp = section.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('Trang ')
    set_run_font(r, 9)
    add_page_number(fp)


doc = Document()
# style mặc định
style = doc.styles['Normal']
style.font.name = FONT
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
set_section_a4(doc.sections[0])
header_footer(doc.sections[0])

# ============================== BÌA ==============================
for _ in range(3):
    p('', space_after=0)
p('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM (HUTECH)', 14, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('KHOA CÔNG NGHỆ THÔNG TIN', 13, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('', space_after=0)
p('BÁO CÁO THỰC HÀNH', 18, True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18)
p('PHÂN TÍCH THIẾT KẾ HƯỚNG ĐỐI TƯỢNG', 14, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('', space_after=0)
p('BÀI 7', 22, True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24)
p('QUẢN LÝ NHÂN SỰ', 22, True, align=WD_ALIGN_PARAGRAPH.CENTER)
p('CÔNG TY SƠN ĐÔNG Á', 15, True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
# đường kẻ ngang
par = doc.add_paragraph()
par.alignment = WD_ALIGN_PARAGRAPH.CENTER
pPr = par._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '12')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '1F3B63')
pBdr.append(bottom)
pPr.append(pBdr)
p('', space_after=0)
info = [
    ('Họ và tên:', 'Lê Minh Giang'),
    ('MSSV:', '1250080043'),
    ('Lớp:', '12_DH_CNPM1'),
    ('Nhóm TH:', '................................'),
    ('Ngày báo cáo:', '29/10/2026'),
    ('Giảng viên hướng dẫn:', '................................'),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
no_borders(t)
for k, v in info:
    row = t.add_row()
    row.cells[0].width = Inches(2.2)
    row.cells[1].width = Inches(3.5)
    cell_margins(row.cells[0], 60, 60, 40, 40)
    cell_margins(row.cells[1], 60, 60, 40, 40)
    pa = row.cells[0].paragraphs[0]
    pa.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ra = pa.add_run(k)
    set_run_font(ra, 12, True)
    pb = row.cells[1].paragraphs[0]
    rb = pb.add_run(v)
    set_run_font(rb, 12)

doc.add_page_break()

# ============================== MỤC LỤC ==============================
heading('MỤC LỤC', 1)
p('1. Mô tả bài toán', 12, space_after=2)
p('2. Trả lời các câu hỏi', 12, space_after=2)
p('   2.1. Câu 1 - Class, thuộc tính, phương thức, ID', 12, space_after=2)
p('   2.2. Câu 2 - Liên kết và Class Diagram', 12, space_after=2)
p('   2.3. Câu 3 - Actor và Use Case tổng quát', 12, space_after=2)
p('   2.4. Câu 4 - Use Case phân rã', 12, space_after=2)
p('   2.5. Câu 5 - Activity Diagram', 12, space_after=2)
p('   2.6. Câu 6 - Sequence Diagram', 12, space_after=2)
p('3. Thiết kế cơ sở dữ liệu và kiến trúc', 12, space_after=2)
p('4. Hiện thực và kiểm thử', 12, space_after=2)
p('5. Kết luận', 12, space_after=2)
doc.add_page_break()

# ============================== 1. MÔ TẢ ==============================
heading('1. MÔ TẢ BÀI TOÁN', 1)
p('Công ty sơn Đông Á cần tin học hóa việc quản lý nhân sự. Bài toán bắt đầu từ khi ứng viên '
  'được tuyển dụng, trải qua thời gian thử việc, ký hợp đồng chính thức, nhận quyết định chức vụ '
  'và lương, làm việc theo kế hoạch tuần, chấm công, ghi nhận giờ phụ trội, khen thưởng hoặc kỷ luật '
  'và cuối cùng là tính bảng lương tháng. Hệ thống phải lưu được các quyết định, hồ sơ và bảng lương '
  'liên quan, bảo đảm dữ liệu nhất quán giữa hợp đồng, bảo hiểm và tiền lương.', first_line=1.0)
p('Các nghiệp vụ chính cần quản lý:', 12, True, space_before=4)
for s in [
    'Tuyển dụng và thử việc từ 1 đến 60 ngày; quyết định tuyển dụng ghi rõ phòng ban, mức lương thử việc và điều khoản hai bên.',
    'Hợp đồng lao động dài hạn hoặc ngắn hạn. Hợp đồng dài hạn bắt buộc lập hồ sơ BHXH (8%) và BHYT (1,5%); hợp đồng ngắn hạn không trích bảo hiểm.',
    'Quyết định chức vụ, phụ cấp chức vụ; quyết định lương, phụ cấp lương và phụ cấp độc hại.',
    'Kế hoạch làm việc tuần, phân công ca, chấm công hằng ngày và phiếu ghi nhận giờ phụ trội (hệ số 1,5).',
    'Khen thưởng, kỷ luật cho cá nhân hoặc tập thể; chi qua phiếu chi, thu qua phiếu thu hoặc trừ vào lương.',
    'Tính và chốt bảng lương tháng từ lương ngày công, phụ trội, phụ cấp, thưởng, phạt và bảo hiểm.',
]:
    bullet(s, 11.5)
p('Phạm vi bài làm: 18 lớp/bảng nghiệp vụ, 6 sơ đồ UML, ứng dụng C# WinForms .NET Framework 4.7.2 '
  'theo kiến trúc UI → Service → Data và bộ kiểm thử 65 case (49 offline + 16 tích hợp CSDL).', 11.5, italic=True)

# ============================== 2. CÂU HỎI ==============================
heading('2. TRẢ LỜI CÁC CÂU HỎI', 1)
heading('2.1. Câu 1 - Class, thuộc tính, phương thức, ID', 2)
p('Em xác định 18 lớp tương ứng với 18 bảng nghiệp vụ. Mỗi lớp có khóa chính đặt tên bắt đầu bằng '
  '“Ma”, khóa ngoại tham chiếu đến lớp khác và 2 phương thức chính. Bảng dưới đây tóm tắt các lớp; '
  'sơ đồ lớp đầy đủ nằm ở Câu 2.', first_line=1.0)
classes = [
    ('1', 'PhongBan', 'MaPB', 'TenPB, DienThoai, TrangThai', 'ThemPhongBan(), CapNhatPB()', 'Phòng tài chính, kỹ thuật, sản xuất...'),
    ('2', 'NhanVien', 'MaNV', 'HoTen, NgaySinh, CCCD, MaPB, TrangThai', 'TiepNhanThuViec(), CapNhatTrangThai()', 'Thông tin nhân thân và trạng thái làm việc'),
    ('3', 'QuyetDinhTuyenDung', 'MaQDTD', 'SoQD, MaNV, MaPB, NgayBatDau, NgayKetThuc, MucLuongThuViec', 'LapQuyetDinh(), KiemTraThoiHan()', 'Quyết định tuyển dụng và thời gian thử việc'),
    ('4', 'HopDongLaoDong', 'MaHD', 'SoHD, MaNV, MaQDTD, LoaiHD, NgayHieuLuc, NgayHetHan, MucLuong', 'KyHopDong(), KiemTraLoaiHD()', 'Hợp đồng dài hạn hoặc ngắn hạn'),
    ('5', 'BaoHiem', 'MaBH', 'MaNV, MaHD, LoaiBH, MucDong, TyLeNV', 'LapHoSoBaoHiem(), TinhKhauTru()', 'Chỉ lập cho hợp đồng dài hạn'),
    ('6', 'QuyetDinhChucVu', 'MaQDCV', 'SoQD, MaNV, MaPB, ChucVu, PhuCapChucVu', 'LapQDCV(), BoNhiem()', 'Bổ nhiệm hoặc điều chuyển chức vụ'),
    ('7', 'QuyetDinhLuong', 'MaQDL', 'SoQD, MaNV, MucLuongCB, PhuCapLuong, PhuCapDocHai', 'LapQDLuong(), TangLuong()', 'Mức lương cơ bản và các phụ cấp'),
    ('8', 'KeHoachTuan', 'MaKH', 'MaPB, TuanSo, Nam, NgayLap, TrangThai', 'LapKeHoach(), DuyetKeHoach()', 'Kế hoạch phân công theo tuần'),
    ('9', 'ChiTietKeHoach', 'MaCTKH', 'MaKH, MaNV, Ngay, CaLam, LaPhuTroi', 'ThemPhanCong(), KiemTraTrung()', 'Từng dòng phân công ca cho nhân viên'),
    ('10', 'ChamCong', 'MaChamCong', 'MaNV, MaCTKH, Ngay, GioVao, GioRa, TrangThai', 'GhiNhanChamCong(), TinhGioLam()', 'Chấm công theo ca và ngày'),
    ('11', 'PhieuPhuTroi', 'MaPhieuPT', 'SoPhieu, NgayPhieu, HinhThuc, HeSo, TrangThai', 'LapPhieuPhuTroi(), DuyetPhieu()', 'Phiếu ghi nhận giờ phụ trội, hệ số 1,5'),
    ('12', 'ChiTietPhuTroi', 'MaCTPT', 'MaPhieuPT, MaNV, SoGio, HeSo', 'ThemChiTiet(), TinhTienPhuTroi()', 'Chi tiết giờ phụ trội của từng nhân viên'),
    ('13', 'QuyetDinhKhenThuong', 'MaQDKhenThuong', 'SoQD, DoiTuong, MaNV, MaPB, SoTien, NguonChiTra', 'LapQDKhenThuong(), TaoPhieuChi()', 'Khen thưởng cá nhân hoặc tập thể'),
    ('14', 'PhieuChiKhenThuong', 'MaPhieuChi', 'SoPhieu, MaQDKhenThuong, NgayChi, SoTien', 'LapPhieuChi(), InPhieuChi()', 'Phiếu chi tiền thưởng'),
    ('15', 'QuyetDinhKyLuat', 'MaQDKyLuat', 'SoQD, DoiTuong, MaNV, MaPB, SoTienPhat, NguonThu', 'LapQDKyLuat(), TaoPhieuThu()', 'Kỷ luật cá nhân hoặc tập thể'),
    ('16', 'PhieuThuNopPhat', 'MaPhieuThu', 'SoPhieu, MaQDKyLuat, NgayThu, SoTien', 'LapPhieuThu(), InPhieuThu()', 'Phiếu thu tiền phạt'),
    ('17', 'BangLuong', 'MaBangLuong', 'Thang, Nam, NgayLap, TongChiPhi, TrangThai', 'TinhBangLuong(), ChotBangLuong()', 'Bảng lương tổng hợp theo tháng'),
    ('18', 'ChiTietBangLuong', 'MaCTBL', 'MaBangLuong, MaNV, LuongCoBan, SoNgayCong, TienPhuTroi, TienThuong, TienKyLuat, BHXH, BHYT, ThucLinh', 'TinhLuong(), TinhThucLinh()', 'Chi tiết lương của từng nhân viên'),
]
add_table(['STT', 'Class', 'ID (PK)', 'Thuộc tính chính', 'Phương thức chính', 'Mô tả'],
          classes, [0.35, 1.45, 1.0, 1.75, 1.55, 0.9], font_size=8.2,
          align_center=[0, 2])
caption('Bảng 2.1 - Danh sách 18 lớp nghiệp vụ và ID tương ứng')

heading('2.2. Câu 2 - Liên kết và Class Diagram', 2)
p('Các lớp liên kết chủ yếu theo quan hệ 1 - 0..* (một phòng ban có nhiều nhân viên, một nhân viên có '
  'nhiều quyết định). Quan hệ 1 - 1..* thể hiện phần chi tiết bắt buộc như kế hoạch tuần phải có ít nhất '
  'một phân công, phiếu phụ trội phải có ít nhất một chi tiết, bảng lương phải có ít nhất một dòng lương. '
  'Quan hệ 1 - 0..1 thể hiện phần mở rộng như chấm công chỉ có khi có phân công, phiếu chi hoặc phiếu thu '
  'chỉ có khi chi/thu tiền.', first_line=1.0)
assoc = [
    ('1', 'PhongBan → NhanVien', '1 - 0..*', 'Mỗi nhân viên thuộc một phòng ban; một phòng ban có nhiều nhân viên.'),
    ('2', 'NhanVien → QuyetDinhTuyenDung', '1 - 0..*', 'Một nhân viên có thể có nhiều quyết định tuyển dụng/thử việc.'),
    ('3', 'NhanVien → HopDongLaoDong', '1 - 0..*', 'Một nhân viên có thể ký nhiều hợp đồng theo thời gian.'),
    ('4', 'HopDongLaoDong → BaoHiem', '1 - 0..*', 'Hợp đồng dài hạn có hồ sơ BHXH, BHYT; ngắn hạn không có.'),
    ('5', 'NhanVien → QuyetDinhChucVu', '1 - 0..*', 'Một nhân viên có nhiều quyết định chức vụ.'),
    ('6', 'NhanVien → QuyetDinhLuong', '1 - 0..*', 'Một nhân viên có nhiều quyết định lương.'),
    ('7', 'PhongBan → KeHoachTuan', '1 - 0..*', 'Mỗi phòng ban lập kế hoạch tuần cho nhân viên.'),
    ('8', 'KeHoachTuan → ChiTietKeHoach', '1 - 1..*', 'Một kế hoạch phải có ít nhất một dòng phân công.'),
    ('9', 'ChiTietKeHoach → ChamCong', '1 - 0..1', 'Một phân công có thể có một bản ghi chấm công.'),
    ('10', 'PhieuPhuTroi → ChiTietPhuTroi', '1 - 1..*', 'Một phiếu phụ trội phải có ít nhất một chi tiết.'),
    ('11', 'QuyetDinhKhenThuong → PhieuChiKhenThuong', '1 - 0..1', 'Quyết định thưởng có thể chi bằng phiếu chi.'),
    ('12', 'QuyetDinhKyLuat → PhieuThuNopPhat', '1 - 0..1', 'Quyết định phạt có thể thu bằng phiếu thu.'),
    ('13', 'BangLuong → ChiTietBangLuong', '1 - 1..*', 'Một bảng lương phải có ít nhất một dòng chi tiết.'),
    ('14', 'NhanVien → QuyetDinhKhenThuong', '1 - 0..*', 'Khen thưởng cá nhân gắn với MaNV.'),
    ('15', 'PhongBan → QuyetDinhKhenThuong', '1 - 0..*', 'Khen thưởng tập thể gắn với MaPB.'),
    ('16', 'NhanVien → ChiTietBangLuong', '1 - 0..*', 'Mỗi dòng lương tháng gắn với một nhân viên.'),
]
add_table(['STT', 'Quan hệ', 'Bản số', 'Ý nghĩa'], assoc,
          [0.4, 2.15, 0.75, 3.7], font_size=8.5, align_center=[0, 2])
caption('Bảng 2.2 - Các quan hệ chính giữa các lớp')

# A3 portrait cho class diagram
sec = doc.add_section(WD_SECTION.NEW_PAGE)
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width = Inches(11.69)
sec.page_height = Inches(16.54)
sec.top_margin = Inches(0.45)
sec.bottom_margin = Inches(0.4)
sec.left_margin = Inches(0.5)
sec.right_margin = Inches(0.5)
header_footer(sec)
add_picture_center(os.path.join(DIAG, '1_ClassDiagram.png'), 8.3)
caption('Hình 2.1 - Sơ đồ lớp giai đoạn thiết kế (18 lớp, ID, khóa ngoại, phương thức). '
        'Nguồn draw.io: diagrams/1_ClassDiagram.drawio.')

# trở lại A4
sec = doc.add_section(WD_SECTION.NEW_PAGE)
set_section_a4(sec)
header_footer(sec)

heading('2.3. Câu 3 - Actor và Use Case tổng quát', 2)
p('Em xác định 6 actor tham gia hệ thống. Mỗi actor được gắn với các use case đúng với vai trò nghiệp vụ; '
  'các use case phụ như lập hồ sơ bảo hiểm hay lập phiếu chi được mô tả là phần mở rộng của use case chính.',
  first_line=1.0)
actors = [
    ('1', 'Nhân viên / Ứng viên', 'Nộp hồ sơ, thử việc, ký hợp đồng, xem thông tin lương.'),
    ('2', 'Trưởng phòng', 'Lập kế hoạch tuần, phân công ca, đánh giá thử việc, đề xuất khen thưởng/kỷ luật.'),
    ('3', 'Ban chấm công', 'Chấm công hằng ngày, lập phiếu ghi nhận giờ phụ trội.'),
    ('4', 'NV HCNS', 'Tiếp nhận hồ sơ, lập quyết định, lập hồ sơ BHXH/BHYT, quản lý khen thưởng/kỷ luật.'),
    ('5', 'Giám đốc', 'Ký quyết định, duyệt khen thưởng/kỷ luật, duyệt và chốt bảng lương.'),
    ('6', 'Kế toán', 'Tính lương, chi thưởng, thu phạt, lập báo cáo nhân sự và tiền lương.'),
]
add_table(['STT', 'Actor', 'Vai trò chính'], actors, [0.4, 1.65, 4.9], font_size=9, align_center=[0])
caption('Bảng 2.3 - Các actor của hệ thống')

ucs = [
    ('UC01', 'Tiếp nhận hồ sơ và thử việc', 'Nhân viên, NV HCNS'),
    ('UC02', 'Ký hợp đồng lao động chính thức', 'Nhân viên, NV HCNS'),
    ('UC03', 'Quyết định chức vụ và phụ cấp', 'Trưởng phòng, Giám đốc'),
    ('UC04', 'Quyết định lương và phụ cấp độc hại', 'Trưởng phòng, Giám đốc'),
    ('UC05', 'Lập hồ sơ BHXH - BHYT', 'NV HCNS'),
    ('UC06', 'Lập kế hoạch tuần và ca làm việc', 'Trưởng phòng'),
    ('UC07', 'Chấm công hằng ngày theo ca', 'Ban chấm công'),
    ('UC08', 'Ghi nhận giờ phụ trội (x1,5)', 'Ban chấm công'),
    ('UC09', 'Khen thưởng và lập phiếu chi', 'Giám đốc, NV HCNS'),
    ('UC10', 'Kỷ luật và thu nộp phạt', 'Giám đốc, NV HCNS'),
    ('UC11', 'Tính và phát hành bảng lương tháng', 'Giám đốc, Kế toán'),
    ('UC12', 'Báo cáo thống kê nhân sự và tiền lương', 'Kế toán'),
]
add_table(['Mã', 'Use case', 'Actor chính'], ucs, [0.65, 3.6, 2.7], font_size=9, align_center=[0])
caption('Bảng 2.4 - Các use case tổng quát')
add_picture_center(os.path.join(DIAG, '2_UseCase_TongQuat.png'), 6.8)
caption('Hình 2.2 - Sơ đồ Use Case tổng quát')

heading('2.4. Câu 4 - Use Case phân rã', 2)
p('Hoạt động từ lúc tiếp nhận thử việc đến khi ký hợp đồng chính thức được em phân rã thành 9 use case con. '
  'Trình tự nghiệp vụ đi từ tiếp nhận hồ sơ, lập quyết định thử việc, phân công và đánh giá, lập quyết định '
  'chức vụ/lương, ký hợp đồng, lập hồ sơ bảo hiểm (nếu hợp đồng dài hạn) và cập nhật trạng thái nhân viên.',
  first_line=1.0)
uc4 = [
    ('UC4.1', 'Tiếp nhận hồ sơ', 'Ứng viên/Nhân viên, HCNS', 'Kiểm tra CCCD 12 số, tuổi ≥ 18, nhu cầu phòng ban.'),
    ('UC4.2', 'Lập QĐ tuyển dụng và thử việc', 'HCNS, Giám đốc', 'Thời gian thử việc 1-60 ngày, mức lương thử việc, điều khoản.'),
    ('UC4.3', 'Phân công thử việc theo ca', 'Trưởng phòng', 'Gán ca làm việc và người hướng dẫn trong kế hoạch tuần.'),
    ('UC4.4', 'Theo dõi và đánh giá thử việc', 'Trưởng phòng, Nhân viên', 'Theo dõi chấm công, đánh giá kết quả cuối thời gian thử việc.'),
    ('UC4.5', 'Lập QĐ chức vụ và phụ cấp', 'HCNS, Giám đốc', 'Bổ nhiệm hoặc điều chuyển, kèm phụ cấp chức vụ nếu có.'),
    ('UC4.6', 'Lập QĐ lương và phụ cấp độc hại', 'HCNS, Giám đốc', 'Mức lương cơ bản, phụ cấp lương, phụ cấp độc hại.'),
    ('UC4.7', 'Ký hợp đồng chính thức', 'Nhân viên, Giám đốc', 'Chọn loại hợp đồng dài hạn hoặc ngắn hạn.'),
    ('UC4.8', 'Lập hồ sơ BHXH - BHYT', 'HCNS', 'Chỉ thực hiện với hợp đồng dài hạn; trích 8% BHXH, 1,5% BHYT.'),
    ('UC4.9', 'Cập nhật trạng thái và lưu hồ sơ', 'HCNS', 'Chuyển nhân viên sang trạng thái “Đang làm”, lưu hồ sơ.'),
]
add_table(['Mã', 'Use case con', 'Actor', 'Nội dung chính'], uc4,
          [0.6, 1.85, 1.55, 3.0], font_size=8.5, align_center=[0])
caption('Bảng 2.5 - Phân rã use case từ thử việc đến ký hợp đồng')
add_picture_center(os.path.join(DIAG, '3_UseCase_PhanRa_ThuViec_HopDong.png'), 6.8)
caption('Hình 2.3 - Sơ đồ Use Case phân rã')

heading('2.5. Câu 5 - Activity Diagram', 2)
p('Sơ đồ hoạt động mô tả tuần tự các bước từ nộp hồ sơ đến ký hợp đồng chính thức. Có ba nút quyết định: '
  'hồ sơ hợp lệ, đạt yêu cầu thử việc và loại hợp đồng. Nếu hồ sơ không hợp lệ thì quay lại bước nộp hồ sơ; '
  'nếu không đạt thử việc thì chấm dứt và thanh toán lương thử việc; nếu hợp đồng dài hạn thì lập thêm hồ sơ '
  'BHXH/BHYT trước khi cập nhật trạng thái.', first_line=1.0)
add_picture_center(os.path.join(DIAG, '4_Activity_ThuViec_HopDong.png'), 5.8)
caption('Hình 2.4 - Activity Diagram từ thử việc đến ký hợp đồng chính thức')

heading('2.6. Câu 6 - Sequence Diagram', 2)
p('Sơ đồ tuần tự cho nghiệp vụ “Lập quyết định khen thưởng” có 4 lifeline: Người dùng, FrmKhenThuong (UI), '
  'KhenThuongService (Service) và SQL Server (Data). Service kiểm tra dữ liệu, mở transaction, insert quyết định, '
  'tự động insert phiếu chi nếu nguồn chi là “Phiếu chi”, commit hoặc rollback và trả ServiceResult về cho Form.',
  first_line=1.0)
add_picture_center(os.path.join(DIAG, '5_Sequence_LapQuyetDinhKhenThuong.png'), 6.8)
caption('Hình 2.5 - Sequence Diagram nghiệp vụ lập quyết định khen thưởng')
p('Ghi chú: nếu nguồn chi là “Lương” thì bỏ qua bước tạo phiếu chi; khoản thưởng được lưu để tổng hợp vào cột '
  'TiềnThưởng của bảng lương tháng.', 10, italic=True)

# ============================== 3. CSDL & KIẾN TRÚC ==============================
heading('3. THIẾT KẾ CƠ SỞ DỮ LIỆU VÀ KIẾN TRÚC', 1)
heading('3.1. Kiến trúc 3 lớp', 2)
p('Ứng dụng được tổ chức theo kiến trúc UI → Service → Data. Form chỉ gọi Service, Service kiểm tra nghiệp vụ '
  'và gọi lớp Db; lớp Db dùng ADO.NET với tham số @parameter để tránh SQL injection. Form không đặt câu SQL trực tiếp.',
  first_line=1.0)
add_picture_center(os.path.join(DIAG, '6_KienTruc_3Lop.png'), 6.8)
caption('Hình 3.1 - Kiến trúc 3 lớp của ứng dụng')

heading('3.2. Các bảng trong CSDL QuanLyNhanSu', 2)
p('CSDL gồm 18 bảng, đặt tên trùng với 18 lớp nghiệp vụ. Khóa chính dùng IDENTITY, khóa ngoại có ràng buộc, '
  'các cột tiền dùng DECIMAL(18,0) theo đơn vị VND.', first_line=1.0)
tables = [
    ('PhongBan', 'Phòng ban', 'NhanVien', 'Nhân viên'),
    ('QuyetDinhTuyenDung', 'QĐ tuyển dụng/thử việc', 'HopDongLaoDong', 'Hợp đồng lao động'),
    ('BaoHiem', 'Hồ sơ BHXH/BHYT', 'QuyetDinhChucVu', 'QĐ chức vụ/phụ cấp'),
    ('QuyetDinhLuong', 'QĐ lương/phụ cấp', 'KeHoachTuan', 'Kế hoạch tuần'),
    ('ChiTietKeHoach', 'Chi tiết phân công ca', 'ChamCong', 'Chấm công'),
    ('PhieuPhuTroi', 'Phiếu phụ trội', 'ChiTietPhuTroi', 'Chi tiết phụ trội'),
    ('QuyetDinhKhenThuong', 'QĐ khen thưởng', 'PhieuChiKhenThuong', 'Phiếu chi thưởng'),
    ('QuyetDinhKyLuat', 'QĐ kỷ luật', 'PhieuThuNopPhat', 'Phiếu thu phạt'),
    ('BangLuong', 'Bảng lương tháng', 'ChiTietBangLuong', 'Chi tiết bảng lương'),
]
add_table(['Bảng', 'Nội dung', 'Bảng', 'Nội dung'], tables,
          [1.45, 2.0, 1.45, 2.0], font_size=8.5)
caption('Bảng 3.1 - Danh mục 18 bảng trong CSDL')

heading('3.3. Quy tắc tính lương', 2)
p('Tháng chuẩn có 26 ngày công, mỗi ngày 8 giờ; giờ phụ trội có hệ số 1,5. Công thức em áp dụng:', first_line=1.0)
for s in [
    'Lương ngày công = MucLuongCB / 26 × min(26, số ngày công).',
    'Tiền phụ trội = MucLuongCB / 26 / 8 × 1,5 × số giờ phụ trội.',
    'Thực lĩnh = lương ngày công + phụ trội + phụ cấp chức vụ + phụ cấp lương + phụ cấp độc hại + thưởng − phạt − BHXH − BHYT.',
    'BHXH = mức đóng × 8%; BHYT = mức đóng × 1,5%; chỉ trích với hợp đồng dài hạn.',
]:
    bullet(s, 11.5)
p('Kết quả chạy thử với dữ liệu mẫu tháng 10/2026:', 11.5, True, space_before=4)
pay = [
    ('NV00001', 'Nguyễn Văn An', '15.000.000', '26', '8', '1.500.000', '2.000.000', '0', '1.200.000', '225.000', '19.440.385'),
    ('NV00002', 'Trần Thị Bình', '12.000.000', '26', '0', '0', '0', '0', '960.000', '180.000', '12.360.000'),
    ('NV00003', 'Lê Hoàng Cường', '9.000.000', '26', '10', '0', '3.000.000', '0', '720.000', '135.000', '14.849.038'),
    ('NV00004', 'Phạm Thu Dung', '18.000.000', '26', '0', '0', '0', '0', '1.440.000', '270.000', '19.790.000'),
    ('NV00005', 'Võ Minh Đức', '8.000.000', '26', '8', '0', '3.000.000', '500.000', '0', '0', '11.546.154'),
    ('NV00006', 'Đặng Thị Hoa', '10.000.000', '26', '0', '0', '0', '0', '800.000', '150.000', '10.050.000'),
]
add_table(['Mã NV', 'Họ tên', 'Lương CB', 'Ngày công', 'Giờ PT', 'PC chức vụ', 'Thưởng', 'Phạt', 'BHXH', 'BHYT', 'Thực lĩnh'],
          pay, [0.7, 1.15, 0.85, 0.55, 0.45, 0.7, 0.65, 0.5, 0.65, 0.55, 0.85], font_size=7.4,
          align_center=[0, 2, 3, 4, 5, 6, 7, 8, 9, 10])
caption('Bảng 3.2 - Kết quả tính lương tháng 10/2026 (đơn vị: VND)')
p('Tổng thực lĩnh 6 nhân viên: 88.035.577 VND.', 11.5, True, align=WD_ALIGN_PARAGRAPH.RIGHT)

heading('3.4. Trigger và stored procedure', 2)
for s in [
    'TR_BaoHiem_ChiDaiHan: chặn insert hồ sơ bảo hiểm cho hợp đồng ngắn hạn.',
    'TR_ChiTietBangLuong_Khoa: chặn sửa chi tiết của bảng lương đã chốt.',
    'sp_TiepNhanThuViec: tiếp nhận nhân viên thử việc trong transaction.',
    'sp_KyHopDongChinhThuc: ký hợp đồng, cập nhật trạng thái và lập bảo hiểm nếu là hợp đồng dài hạn.',
    'sp_ChotBangLuong: chốt bảng lương khi đã có chi tiết.',
]:
    bullet(s, 11.5)

# ============================== 4. HIỆN THỰC & KIỂM THỬ ==============================
heading('4. HIỆN THỰC VÀ KIỂM THỬ', 1)
heading('4.1. Công nghệ và các form chính', 2)
p('Ứng dụng viết bằng C# WinForms trên .NET Framework 4.7.2, kết nối SQL Server qua ADO.NET. Solution đặt tại '
  'QuanLyNhanSu/QuanLyNhanSu.slnx; CSDL tạo bằng hai script SQL/01_CreateTables.sql và SQL/02_SeedData.sql.',
  first_line=1.0)
forms = [
    ('FrmTiepNhan', 'Tiếp nhận nhân viên thử việc', 'TuyenDungService'),
    ('FrmHopDong', 'Ký hợp đồng và lập hồ sơ bảo hiểm', 'HopDongService'),
    ('FrmQuyetDinh', 'Quyết định chức vụ và lương', 'QuyetDinhService'),
    ('FrmKeHoach', 'Lập kế hoạch tuần và phân công ca', 'KeHoachService'),
    ('FrmChamCong', 'Chấm công theo kế hoạch', 'ChamCongService'),
    ('FrmPhuTroi', 'Lập phiếu ghi nhận giờ phụ trội', 'PhuTroiService'),
    ('FrmKhenThuong', 'Khen thưởng và kỷ luật', 'KhenThuongService'),
    ('FrmLuong', 'Tính và chốt bảng lương tháng', 'LuongService'),
    ('FrmDanhSach', 'Xem danh sách nhân viên và hợp đồng', 'Db'),
    ('FrmKetNoi', 'Cấu hình kết nối CSDL', 'Db'),
]
add_table(['Form', 'Chức năng', 'Service xử lý'], forms, [1.35, 3.35, 2.25], font_size=9)
caption('Bảng 4.1 - Các form chính và service tương ứng')

heading('4.2. Kết quả kiểm thử', 2)
p('Em kiểm thử ở hai mức: 49 case offline (không cần CSDL) và 16 case tích hợp trên CSDL thật. Kết quả chạy '
  'trên form FrmMain: 49/49 PASS và 16/16 PASS, không có case lỗi.', first_line=1.0)
groups = [
    ('TD01 - TD09', '9', 'Tiếp nhận, tuổi, CCCD, thời gian thử việc, ngày ban hành.'),
    ('HD01 - HD08', '8', 'Loại hợp đồng, ngày hiệu lực, bảo hiểm dài hạn/ngắn hạn.'),
    ('KH01 - KH05', '5', 'Tuần làm việc, ngày ngoài tuần, phân công trùng, ca làm.'),
    ('CC01 - CC06', '6', 'Giờ vào/ra, vắng mặt, tính số giờ làm.'),
    ('PT01 - PT06', '6', 'Số giờ phụ trội, hệ số 1,5, trùng nhân viên, tính tiền.'),
    ('KT01 - KT05, KL01 - KL02', '7', 'Khen thưởng/kỷ luật cá nhân, tập thể, chia đều, nguồn chi.'),
    ('L01 - L08', '8', 'Lương ngày công, phụ trội, BHXH, BHYT, thực lĩnh, không âm.'),
    ('DB01 - DB16', '16', 'Kết nối, đủ 18 bảng, seed, bảng lương, trigger, transaction.'),
]
add_table(['Nhóm test', 'Số case', 'Nội dung kiểm tra'], groups, [1.75, 0.7, 4.5], font_size=8.8, align_center=[1])
caption('Bảng 4.2 - Tổng hợp 65 test case')
cases = [
    ('TD01', 'Tiếp nhận hồ sơ hợp lệ', 'Validate trả về null', 'PASS'),
    ('TD03', 'Ứng viên dưới 18 tuổi', 'Bị chặn', 'PASS'),
    ('TD05', 'Thử việc 70 ngày', 'Bị chặn (quá 60 ngày)', 'PASS'),
    ('HD03', 'Hợp đồng dài hạn có ngày hết hạn', 'Bị chặn', 'PASS'),
    ('KH01', 'Tuần 42/2026', 'Bắt đầu 12/10/2026', 'PASS'),
    ('CC05', 'Giờ vào 08:00, giờ ra 17:00', '8 giờ làm', 'PASS'),
    ('PT06', '8 giờ phụ trội, lương 15 triệu', '865.385 VND', 'PASS'),
    ('KT05', 'Chia thưởng tập thể 6 triệu cho 2 NV', '3.000.000 VND/người', 'PASS'),
    ('L07', 'Thực lĩnh NV1 tháng 10/2026', '19.440.385 VND', 'PASS'),
    ('DB04', 'Tính bảng lương tháng 10/2026', '6 nhân viên, tổng 88.035.577', 'PASS'),
    ('DB09', 'Trigger bảo hiểm hợp đồng ngắn hạn', 'Chặn insert (lỗi 50011)', 'PASS'),
    ('DB16', 'Chốt bảng lương và sửa chi tiết', 'Chặn sửa (lỗi 50010)', 'PASS'),
]
add_table(['Mã', 'Tình huống', 'Kết quả mong đợi', 'Kết quả'], cases,
          [0.6, 2.75, 2.9, 0.7], font_size=8.5, align_center=[0, 3])
caption('Bảng 4.3 - Một số test case tiêu biểu')

p('Ảnh chụp giao diện và kết quả kiểm thử:', 11.5, True, space_before=4)
# grid 2 cột, 6 hàng
shots = [
    ('Screenshot 2026-10-29 105530.png', 'Hình 4.1 - FrmMain: menu nghiệp vụ'),
    ('Screenshot 2026-10-29 105536.png', 'Hình 4.2 - FrmTiepNhan: tiếp nhận thử việc'),
    ('Screenshot 2026-10-29 105542.png', 'Hình 4.3 - FrmHopDong: ký hợp đồng và bảo hiểm'),
    ('Screenshot 2026-10-29 105601.png', 'Hình 4.4 - FrmKeHoach: kế hoạch tuần'),
    ('Screenshot 2026-10-29 105607.png', 'Hình 4.5 - FrmChamCong: chấm công theo ca'),
    ('Screenshot 2026-10-29 105612.png', 'Hình 4.6 - FrmPhuTroi: phiếu phụ trội'),
    ('Screenshot 2026-10-29 105617.png', 'Hình 4.7 - FrmKhenThuong: khen thưởng/kỷ luật'),
    ('Screenshot 2026-10-29 105622.png', 'Hình 4.8 - FrmLuong: tính và chốt bảng lương'),
    ('Screenshot 2026-10-29 105631.png', 'Hình 4.9 - FrmKetNoi: cấu hình kết nối CSDL'),
    ('Screenshot 2026-10-29 105640.png', 'Hình 4.10 - Kết quả 49/49 test offline'),
    ('Screenshot 2026-10-29 105645.png', 'Hình 4.11 - Kết quả 16/16 test CSDL'),
    ('Screenshot 2026-10-29 105650.png', 'Hình 4.12 - FrmDanhSach: danh sách nhân viên'),
]
grid = doc.add_table(rows=0, cols=2)
grid.alignment = WD_TABLE_ALIGNMENT.CENTER
grid.autofit = False
no_borders(grid)
for i in range(0, len(shots), 2):
    row = grid.add_row()
    for j in range(2):
        cell = row.cells[j]
        cell.width = Inches(3.35)
        cell_margins(cell, 30, 30, 30, 30)
        if i + j < len(shots):
            fn, cap = shots[i + j]
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.add_run().add_picture(os.path.join(WORD, fn), width=Inches(3.25))
            cp = cell.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(4)
            r = cp.add_run(cap)
            set_run_font(r, 8.5, italic=True)

doc.add_page_break()
heading('4.3. Truy vết yêu cầu - UML - Form/Service - CSDL - Test', 2)
trace = [
    ('Tuyển dụng, thử việc', 'UC01, UC4.1-4.4', 'FrmTiepNhan / TuyenDungService', 'NhanVien, QuyetDinhTuyenDung', 'TD01-TD09, DB11'),
    ('Hợp đồng lao động', 'UC02, UC4.7', 'FrmHopDong / HopDongService', 'HopDongLaoDong', 'HD01-HD08, DB10'),
    ('Bảo hiểm dài hạn', 'UC02, UC4.8', 'FrmHopDong / HopDongService', 'BaoHiem', 'HD07-HD08, DB08-DB09'),
    ('QĐ chức vụ, lương', 'UC03, UC04', 'FrmQuyetDinh / QuyetDinhService', 'QuyetDinhChucVu, QuyetDinhLuong', 'L01-L08'),
    ('Kế hoạch tuần', 'UC06', 'FrmKeHoach / KeHoachService', 'KeHoachTuan, ChiTietKeHoach', 'KH01-KH05, DB12'),
    ('Chấm công', 'UC07', 'FrmChamCong / ChamCongService', 'ChamCong', 'CC01-CC06, DB13'),
    ('Phụ trội', 'UC08', 'FrmPhuTroi / PhuTroiService', 'PhieuPhuTroi, ChiTietPhuTroi', 'PT01-PT06, L04'),
    ('Khen thưởng', 'UC09', 'FrmKhenThuong / KhenThuongService', 'QDKT, PhieuChiKhenThuong', 'KT01-KT05, DB14'),
    ('Kỷ luật', 'UC10', 'FrmKhenThuong / KhenThuongService', 'QDKL, PhieuThuNopPhat', 'KL01-KL02, DB15'),
    ('Bảng lương tháng', 'UC11', 'FrmLuong / LuongService', 'BangLuong, ChiTietBangLuong', 'L01-L08, DB04-DB07, DB16'),
    ('Báo cáo, danh sách', 'UC12', 'FrmDanhSach', 'NhanVien, HopDongLaoDong', 'DB03'),
]
add_table(['Yêu cầu', 'Use case', 'Form / Service', 'Bảng CSDL', 'Test'],
          trace, [1.2, 1.15, 1.75, 1.65, 1.15], font_size=7.8)
caption('Bảng 4.4 - Ma trận truy vết yêu cầu → UML → Form/Service → CSDL → Test')

# ============================== 5. KẾT LUẬN ==============================
heading('5. KẾT LUẬN', 1)
p('Bài làm đã hoàn thành 6 câu hỏi của bài 7: xác định 18 lớp và ID, vẽ sơ đồ lớp có bản số, xác định actor '
  'và 12 use case tổng quát, phân rã 9 use case từ thử việc đến ký hợp đồng, vẽ activity diagram và sequence '
  'diagram cho nghiệp vụ khen thưởng. Phần hiện thực gồm CSDL 18 bảng, ứng dụng WinForms theo kiến trúc 3 lớp '
  'và bộ kiểm thử 65 case; kết quả chạy thực tế là 49/49 case offline và 16/16 case CSDL đều đạt.', first_line=1.0)
p('Hạn chế:', 11.5, True, space_before=4)
for s in ['Chưa có đăng nhập và phân quyền theo vai trò.',
          'Báo cáo mới dừng ở danh sách, chưa xuất Excel/PDF.',
          'Chưa tính thuế thu nhập cá nhân và các khoản giảm trừ khác.']:
    bullet(s, 11)
p('Hướng phát triển:', 11.5, True, space_before=4)
for s in ['Thêm đăng nhập, phân quyền và ghi nhật ký thao tác.',
          'Xuất báo cáo Excel/PDF, gửi phiếu lương qua email.',
          'Bổ sung tính thuế TNCN, bảo hiểm thất nghiệp và các khoản phúc lợi.']:
    bullet(s, 11)

doc.save(OUT)
print('saved', OUT)
